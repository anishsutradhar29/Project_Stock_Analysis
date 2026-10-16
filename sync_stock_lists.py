"""
Regenerate the hardcoded stock lists inside SGX_Stock_Analysis.py and NSE_Stock_Analysis.py
from 'stock list.xlsx' - the spreadsheet is the one place to edit which stocks are analysed.

- 'SGX' worksheet -> SGX_INDEX_STOCK_ROWS (SGX's entire default-mode universe)
- 'NSE' worksheet -> NSE_PERSONAL_STOCK_ROWS (personal additions on top of NSE's Top N + NIFTY 50)

Each worksheet has a header row (Index / Inclusion reason, Company, Ticker) followed by one row
per stock, every cell filled in. Every row is validated before anything is written: if any
row is invalid, all problems are printed and neither script is touched.

Only the text between each script's BEGIN/END GENERATED STOCK LIST markers is replaced, and
only if it actually changed. Runs automatically as the first step of update_notebook.py, or
on its own:

    python sync_stock_lists.py [--xlsx "stock list.xlsx"]

Note: this reads the last *saved* version of the spreadsheet - unsaved edits still open in
Excel are not picked up.
"""

import argparse
import ast
import json
import os
import sys

# Resolve paths relative to this script's own location (not a hardcoded absolute
# path) so this works unchanged after copying the project to another machine/OS.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(SCRIPT_DIR, "stock list.xlsx")

BEGIN_MARKER = "# >>> BEGIN GENERATED STOCK LIST (sync_stock_lists.py) - edit stock list.xlsx, not this block\n"
END_MARKER = "# <<< END GENERATED STOCK LIST\n"

# One entry per worksheet: which script/constant it generates, and what a valid row looks like.
SHEETS = [
    {
        "sheet": "SGX",
        "script": "SGX_Stock_Analysis.py",
        "constant": "SGX_INDEX_STOCK_ROWS",
        "header": ("Index", "Company", "Ticker"),
        "allowed_first_col": ("STI", "S Cap", "L & M Cap", "Personal"),
        "ticker_suffix": ".SI",
    },
    {
        "sheet": "NSE",
        "script": "NSE_Stock_Analysis.py",
        "constant": "NSE_PERSONAL_STOCK_ROWS",
        "header": ("Inclusion reason", "Company", "Ticker"),
        "allowed_first_col": ("Personal",),
        "ticker_suffix": ".NS",
    },
]


class StockListError(Exception):
    """Raised when the spreadsheet or a script's generated block can't be synced safely."""


def cell_text(value):
    """Normalise one cell to a stripped string, or None if it's empty/whitespace-only."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_sheet_rows(wb, spec):
    """Read and validate one worksheet. Returns (rows, errors) - rows as (col1, company, ticker)."""
    name = spec["sheet"]
    if name not in wb.sheetnames:
        return [], [f"'{name}' worksheet not found (sheets present: {wb.sheetnames})"]

    rows, errors = [], []
    header_seen = False
    seen_tickers = {}
    for excel_row, raw in enumerate(wb[name].iter_rows(values_only=True), start=1):
        # Data doesn't necessarily start in column A, so drop empty cells rather than
        # relying on fixed column positions.
        values = [t for t in (cell_text(c) for c in raw) if t is not None]
        if not values:
            continue
        where = f"{name} row {excel_row}"

        if not header_seen:
            header_seen = True
            if tuple(v.lower() for v in values) != tuple(h.lower() for h in spec["header"]):
                errors.append(f"{where}: expected header {list(spec['header'])}, found {values}")
            continue

        if len(values) != 3:
            errors.append(
                f"{where}: expected 3 filled-in cells ({', '.join(spec['header'])}), "
                f"found {len(values)}: {values}"
            )
            continue

        first, company, ticker = values
        if first not in spec["allowed_first_col"]:
            errors.append(
                f"{where}: '{spec['header'][0]}' must be one of {list(spec['allowed_first_col'])}, found '{first}'"
            )
        if not ticker.upper().endswith(spec["ticker_suffix"]) or len(ticker) <= len(spec["ticker_suffix"]):
            errors.append(f"{where}: ticker '{ticker}' should look like CODE{spec['ticker_suffix']}")
        ticker = ticker.upper()
        if ticker in seen_tickers:
            errors.append(f"{where}: duplicate ticker '{ticker}' (already on row {seen_tickers[ticker]})")
        seen_tickers.setdefault(ticker, excel_row)
        rows.append((first, company, ticker))

    if not header_seen:
        errors.append(f"'{name}' worksheet is empty")
    elif not rows and not errors:
        errors.append(f"'{name}' worksheet has a header but no stock rows")
    return rows, errors


def render_rows(constant, rows):
    """Render rows as a Python list literal. json.dumps gives double-quoted, escaped strings
    that are also valid Python literals (handles apostrophes, e.g. "Kwality Wall's")."""
    lines = [f"{constant} = ["]
    for row in rows:
        lines.append("    (" + ", ".join(json.dumps(v, ensure_ascii=False) for v in row) + "),")
    lines.append("]")
    text = "\n".join(lines) + "\n"
    # Belt and braces: the generated text must parse back to exactly these rows.
    if ast.literal_eval(text.split("=", 1)[1].strip()) != [tuple(r) for r in rows]:
        raise StockListError(f"internal error: generated {constant} doesn't round-trip")
    return text


def split_generated_block(source, script):
    """Split a script's source into (before, block, after) around the generated markers."""
    if source.count(BEGIN_MARKER) != 1 or source.count(END_MARKER) != 1:
        raise StockListError(
            f"{script}: expected exactly one BEGIN and one END GENERATED STOCK LIST marker line"
        )
    start = source.index(BEGIN_MARKER) + len(BEGIN_MARKER)
    end = source.index(END_MARKER)
    if end < start:
        raise StockListError(f"{script}: END marker comes before BEGIN marker")
    return source[:start], source[start:end], source[end:]


def parse_existing_rows(block):
    """Rows currently in a generated block (for the added/removed report). [] if unparseable."""
    try:
        return [tuple(r) for r in ast.literal_eval(block.split("=", 1)[1].strip())]
    except Exception:
        return []


def sync_stock_lists(xlsx_path=DEFAULT_XLSX, script_dir=SCRIPT_DIR):
    """
    Regenerate both scripts' stock lists from the spreadsheet. Raises StockListError (and
    writes nothing) if anything is invalid. Returns True if either script was changed.
    """
    try:
        import openpyxl
    except ImportError:
        raise StockListError("openpyxl isn't installed for this Python - run it with the project's venv")

    if not os.path.exists(xlsx_path):
        raise StockListError(f"spreadsheet not found: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    # Pass 1: read + validate everything, and prepare every new script source in memory.
    errors, pending = [], []
    for spec in SHEETS:
        rows, sheet_errors = read_sheet_rows(wb, spec)
        errors.extend(sheet_errors)
        script_path = os.path.join(script_dir, spec["script"])
        with open(script_path, "r", encoding="utf-8") as f:
            source = f.read()
        before, old_block, after = split_generated_block(source, spec["script"])
        if not sheet_errors:
            new_block = render_rows(spec["constant"], rows)
            pending.append((spec, script_path, rows, old_block, before + new_block + after, source))
    wb.close()

    if errors:
        raise StockListError(
            f"{len(errors)} problem(s) in '{os.path.basename(xlsx_path)}' - nothing was changed:\n  - "
            + "\n  - ".join(errors)
        )

    # Pass 2: everything is valid - write the scripts that changed and report.
    changed = False
    print(f"Syncing stock lists from '{os.path.basename(xlsx_path)}' (last saved version):")
    for spec, script_path, rows, old_block, new_source, old_source in pending:
        old_rows = parse_existing_rows(old_block)
        old_tickers = {r[2] for r in old_rows}
        new_tickers = {r[2] for r in rows}
        added = [r for r in rows if r[2] not in old_tickers]
        removed = [r for r in old_rows if r[2] not in new_tickers]

        label = f"  {spec['script']} ({spec['constant']}): {len(rows)} stocks"
        if new_source == old_source:
            print(f"{label} - no changes")
            continue
        with open(script_path, "w", encoding="utf-8") as f:
            f.write(new_source)
        changed = True
        if added or removed:
            print(f"{label} (+{len(added)} / -{len(removed)})")
        else:
            print(f"{label} (same stocks; index/name/order details updated)")
        for r in added:
            print(f"      + {r[2]:<16} {r[1]}  [{r[0]}]")
        for r in removed:
            print(f"      - {r[2]:<16} {r[1]}  [{r[0]}]")
    return changed


def main():
    parser = argparse.ArgumentParser(description="Regenerate both scripts' hardcoded stock lists from stock list.xlsx.")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="Path to the stock list workbook (default: 'stock list.xlsx' next to this script).")
    args = parser.parse_args()
    try:
        sync_stock_lists(args.xlsx)
    except StockListError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
